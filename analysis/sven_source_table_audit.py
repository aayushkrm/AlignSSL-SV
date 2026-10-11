"""Read-only, finite audit of the published SVEN figure source archive.

This reproduces absolute-expression correlations. It cannot measure SV effects.
Use the bundled scientific Python runtime; no model or genomic reads required.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
from pathlib import Path
from statistics import mean
import zipfile

SOURCE_MEMBER = "41467_2024_55392_MOESM10_ESM.zip"
WORKBOOK_MEMBER = "Source Data/Source Data.xlsx"
SOURCE_BYTES = 10_105_269
SOURCE_MD5 = "94d0a1d10a542bcef0071baa050ebfad"
MAX_OUTER_BYTES = 32 * 1024 * 1024
MAX_WORKBOOK_BYTES = 16 * 1024 * 1024


def average_ranks(values: list[float]) -> list[float]:
    order = sorted(range(len(values)), key=values.__getitem__)
    ranks = [0.0] * len(values)
    start = 0
    while start < len(order):
        end = start + 1
        while end < len(order) and values[order[end]] == values[order[start]]:
            end += 1
        rank = (start + 1 + end) / 2
        for index in order[start:end]:
            ranks[index] = rank
        start = end
    return ranks


def spearman_correlation(x: list[float], y: list[float]) -> float:
    if len(x) != len(y) or len(x) < 2:
        raise ValueError("Rank correlation requires equal, nontrivial vectors")
    rx, ry = average_ranks(x), average_ranks(y)
    mx, my = mean(rx), mean(ry)
    cross = math.fsum((a - mx) * (b - my) for a, b in zip(rx, ry))
    var_x = math.fsum((a - mx) ** 2 for a in rx)
    var_y = math.fsum((b - my) ** 2 for b in ry)
    if var_x == 0 or var_y == 0:
        raise ValueError("Undefined correlation for constant vector")
    return cross / math.sqrt(var_x * var_y)


def unique_member(archive: zipfile.ZipFile, name: str, limit: int) -> bytes:
    matches = [entry for entry in archive.infolist() if entry.filename == name]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one member: {name}")
    if matches[0].file_size > limit:
        raise ValueError(f"Member exceeds read cap: {name}")
    return archive.read(matches[0])  # ZIP CRC checked by ZipFile.


def absolute_expression_summary(sheet, expected_rows: int = 94_366) -> dict:
    if sheet.max_column != 2:
        raise ValueError(f"Unexpected columns for {sheet.title}")
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if header != ("True value", "Predicted Value"):
        raise ValueError(f"Unexpected header for {sheet.title}: {header}")
    observed, predicted = [], []
    for index, row in enumerate(rows, start=2):
        if len(row) != 2 or any(
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
            for value in row
        ):
            raise ValueError(f"Invalid numeric pair in {sheet.title} row {index}")
        observed.append(float(row[0]))
        predicted.append(float(row[1]))
    if len(observed) != expected_rows:
        raise ValueError(f"Unexpected row count for {sheet.title}: {len(observed)}")
    rho = spearman_correlation(observed, predicted)
    if not math.isfinite(rho):
        raise ValueError("Undefined rank correlation")
    return {
        "sheet": sheet.title,
        "range": f"A1:B{len(observed) + 1}",
        "headers": list(header),
        "data_rows": len(observed),
        "spearman_absolute_log10_expression": rho,
        "complete_numeric_rows": len(observed),
        "missing_rows_dropped": 0,
        "gene_or_sample_identifiers_present": False,
        "matched_reference_predictions_present": False,
        "causal_effect_accuracy": "NOT_IDENTIFIABLE_FROM_THIS_SHEET",
    }


def audit(path: Path) -> dict:
    # Keep pure rank/schema controls usable without the optional workbook reader.
    import openpyxl

    if path.stat().st_size > MAX_OUTER_BYTES:
        raise ValueError("Outer archive exceeds read cap")
    outer_bytes = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(outer_bytes)) as outer:
        source = unique_member(outer, SOURCE_MEMBER, MAX_OUTER_BYTES)
    if len(source) != SOURCE_BYTES or hashlib.md5(source).hexdigest() != SOURCE_MD5:
        raise ValueError("Publisher-listed Source Data size/MD5 mismatch")
    with zipfile.ZipFile(io.BytesIO(source)) as inner:
        workbook_bytes = unique_member(inner, WORKBOOK_MEMBER, MAX_WORKBOOK_BYTES)
        members = [
            {"name": entry.filename, "bytes": entry.file_size}
            for entry in inner.infolist()
            if not entry.is_dir()
        ]
    # No extractall and no write/recalculation/export of the authors' workbook.
    workbook = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes), read_only=True, data_only=False
    )
    try:
        inventory = [
            {"sheet": sheet.title, "declared_rows": sheet.max_row,
             "declared_columns": sheet.max_column}
            for sheet in workbook
        ]
        summaries = [
            absolute_expression_summary(workbook[name])
            for name in ("Figure 3a", "Figure S4")
        ]
    finally:
        workbook.close()
    return {
        "schema": 1,
        "scope": "READ_ONLY_PUBLISHED_FIGURE_REPRODUCTION",
        "source_doi": "10.1038/s41467-024-55392-7",
        "source_url": "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC11685779/supplementaryFiles",
        "outer_archive_bytes": len(outer_bytes),
        "outer_archive_sha256": hashlib.sha256(outer_bytes).hexdigest(),
        "source_member": SOURCE_MEMBER,
        "source_member_bytes": len(source),
        "source_member_md5": hashlib.md5(source).hexdigest(),
        "source_member_sha256": hashlib.sha256(source).hexdigest(),
        "publisher_source_md5_match": True,
        "workbook_member": WORKBOOK_MEMBER,
        "workbook_bytes": len(workbook_bytes),
        "workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "source_members": members,
        "workbook_inventory": inventory,
        "absolute_expression_benchmarks": summaries,
        "claims_not_supported": [
            "Measured SV-induced expression effect accuracy",
            "Gene-balanced performance without a verified identifier crosswalk",
            "Reference-only baseline comparison without reference predictions",
            "Independent-allele sample size from unlabelled plotting rows",
            "A model failure, new method, or publication contribution",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = audit(args.archive)
    payload = json.dumps(result, indent=2, allow_nan=False) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open("x", encoding="utf-8") as handle:
            handle.write(payload)  # Never overwrite a prior result.
    print(payload, end="")


if __name__ == "__main__":
    main()
